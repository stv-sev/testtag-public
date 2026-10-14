"""P1B: liest eine Testtag-Arbeitsmappe nur lesend in die Rohdaten des Publikationsmodells.

Wiederverwendet aus P0 (`excel_probe`): Stabilitätsprüfung und temporäre Kopie (`stable_temp_copy`),
Laden mit openpyxl (`_load`, read-only, gespeicherte Resultate), Prüfsumme (`sha256_file`) und die
Public-Prüfung von B2–B6 (`read_public`).

Zusätzlich wird das XML der verwendeten Blätter geprüft: Formelzellen ohne gespeichertes Resultat und
Excel-Fehlerwerte führen nur dann zu einem Fehler, wenn die Zelle tatsächlich als Quelle gelesen wird.
Kosmetische Reste (z. B. `#REF!` in Hilfsspalten) bleiben dadurch ohne Wirkung.

`Noteneingabe` wird über die Kopfzeilen (Gerät, Linie, Sonderzeichen, Suchbereich, Ergebnisbereich, Feld)
strukturell erkannt, nicht über feste Spaltenpositionen.
"""
from __future__ import annotations

import posixpath
import re
import time
import zipfile
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any, Callable, Iterable, Optional
from xml.etree import ElementTree as ET

from . import excel_probe
from .excel_probe import ValidationError
from .publication_model import (
    ATHLETIK_STATION, BEREICH_BEWEGLICHKEIT, BEREICH_KRAFT, BEREICH_TECHNIK, EINTEILUNG_FELDER, QUELLE_DIREKT,
    QUELLE_STATION, Code,
    DirectField, Issues, LineDefinition, OnsiteBlock, ProgramSource, RawParticipant, ScoreSlot, SourceInfo,
    StartlistEntry, Station, build_state, clean_text,
)

SHEETS = ("Public", "Anmeldung", "Testprogramm", "Startliste", "Zeitplan", "Noteneingabe")
BLOCK_FIELDS = ("id_athlete", "id_rotation", "id_apparatus", "id_position", "id_linie", "SKS", "TAL", "End")
NOTEN_LABELS = ("Gerät", "Linie", "Sonderzeichen", "Suchbereich", "Ergebnissbereich")
ANMELDUNG_HEADERS = ("ID", "Name", "Vorname", "Jg.", "Verein", "AK", "Startgeraet", "akKontr")
MODE_FIELDS = (("A7", "B7", "Kraftmodus", "kraft_modus"), ("A8", "B8", "Beweglichkeitsmodus", "beweglichkeit_modus"))
KRAFT_SHEET = re.compile(r"^9(\d+)")
LINE_NUMBER = re.compile(r"^ç(\d+)§(\d+)\|0$")

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL = re.compile(r"^([A-Z]+)(\d+)$")


def col_letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def _col_index(letters: str) -> int:
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(letters)


# ── XML-Prüfung: fehlende Formelresultate und Fehlerwerte ────────────────────────────────────────────

def _sheet_members(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target") for rel in rels.iter(f"{_PKG_REL_NS}Relationship")}
    members = {}
    for sheet in workbook.iter(f"{_NS}sheet"):
        target = targets.get(sheet.get(f"{_REL_NS}id"), "")
        members[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
    return members


def scan_cell_flags(path: Path, sheet_names: Iterable[str]) -> dict[str, dict[tuple[int, int], str]]:
    """Je Blatt: {(Zeile, Spalte): 'fehlt' | 'fehlerwert'} für Formelzellen ohne Resultat und Fehlerwerte."""
    flags: dict[str, dict[tuple[int, int], str]] = {}
    with zipfile.ZipFile(path) as archive:
        members = _sheet_members(archive)
        for name in sheet_names:
            found: dict[tuple[int, int], str] = {}
            with archive.open(members[name]) as handle:
                for _event, elem in ET.iterparse(handle, events=("end",)):
                    if elem.tag != f"{_NS}c":
                        continue
                    match = _CELL.match(elem.get("r", ""))
                    if match:
                        kind = elem.get("t")
                        formula = elem.find(f"{_NS}f")
                        value = elem.find(f"{_NS}v")
                        text = value.text if value is not None else None
                        coord = (int(match.group(2)), _col_index(match.group(1)))
                        if kind == "e":
                            found[coord] = "fehlerwert"
                        elif formula is not None and not text and kind not in ("str", "inlineStr"):
                            found[coord] = "fehlt"
                    elem.clear()
            flags[name] = found
    return flags


class _Sheet:
    """Gespeicherte Werte eines Blatts; `get` meldet fehlende Formelresultate und Fehlerwerte der gelesenen Zelle."""

    def __init__(self, name: str, ws, flags: dict[tuple[int, int], str], issues: Issues, program: str):
        self.name = name
        self.flags = flags
        self.issues = issues
        self.program = program
        ws.reset_dimensions()
        self.values = {(c.row, c.column): c.value for c in excel_probe._iter_content_cells(ws)}
        self.max_row = max((r for r, _c in self.values), default=0)
        self.max_col = max((c for _r, c in self.values), default=0)
        self._reported: set = set()

    def raw(self, row: int, col: int) -> Any:
        return self.values.get((row, col))

    def get(self, row: int, col: int) -> Any:
        flag = self.flags.get((row, col))
        if flag:
            if (row, col) not in self._reported:
                self._reported.add((row, col))
                where = f"{self.program}: {self.name}!{col_letter(col)}{row}"
                if flag == "fehlt":
                    self.issues.add(Code.FORMELRESULTAT_FEHLT, f"{where}: Formel ohne gespeichertes Resultat "
                                                               "(in Excel neu berechnen und speichern).")
                else:
                    self.issues.add(Code.EXCEL_FEHLERWERT, f"{where}: Excel-Fehlerwert in einer verwendeten Zelle.")
            return None
        return self.values.get((row, col))

    def text(self, row: int, col: int) -> str:
        value = self.raw(row, col)
        return value.strip() if isinstance(value, str) else ""

    def find(self, wanted: str, rows: range, cols: Optional[range] = None) -> Optional[tuple[int, int]]:
        cols = cols or range(1, self.max_col + 1)
        for r in rows:
            for c in cols:
                if self.text(r, c) == wanted:
                    return r, c
        return None


def _header_columns(sheet: _Sheet, row: int, cols: range) -> dict[str, int]:
    """Kopftext -> Spalte; bei mehrfach vorkommendem Kopf gilt das erste Vorkommen (z. B. 'AK' und Tab_AK)."""
    found: dict[str, int] = {}
    for c in cols:
        if sheet.text(row, c):
            found.setdefault(sheet.text(row, c), c)
    return found


def _structure(issues: Issues, program: str, message: str) -> None:
    issues.add(Code.STRUKTUR_UNGUELTIG, f"{program}: {message}")


# ── Blattleser ───────────────────────────────────────────────────────────────────────────────────────

def _read_public(ws, sheet: _Sheet, issues: Issues) -> dict:
    try:
        public = excel_probe.read_public(ws)
    except ValidationError as exc:
        issues.add(Code.PUBLIC_UNGUELTIG, str(exc))
        issues.raise_if_any()
    for row in range(2, 9):
        sheet.get(row, 2)
    for label_cell, value_cell, label, key in MODE_FIELDS:
        row = int(label_cell[1:])
        if sheet.text(row, 1) != label:
            issues.add(Code.PUBLIC_UNGUELTIG, f"Public!{label_cell} muss die Beschriftung '{label}' enthalten.")
        public[key] = sheet.get(int(value_cell[1:]), 2)
    return public


def _read_stations(sheet: _Sheet, issues: Issues, program: str) -> list[Station]:
    head = sheet.find("startgeraet", range(1, 40))
    if head is None:
        _structure(issues, program, "Anmeldung: Tabelle tab_startgeraet (Kopf 'startgeraet') nicht gefunden.")
        return []
    row, col = head
    id_col = next((c for c in range(col, col + 6) if sheet.text(row, c) == "GerID"), None)
    if id_col is None:
        _structure(issues, program, "Anmeldung: Spalte 'GerID' in tab_startgeraet fehlt.")
        return []
    stations = []
    r = row + 1
    while sheet.text(r, col):
        code = sheet.get(r, id_col)
        if isinstance(code, bool) or not isinstance(code, (int, float)) or int(code) != code:
            _structure(issues, program, f"Anmeldung!{col_letter(id_col)}{r}: GerID ist keine ganze Zahl.")
        else:
            stations.append(Station(sheet.text(r, col), int(code)))
        r += 1
    return stations


def _read_participants(sheet: _Sheet, issues: Issues, program: str) -> list[RawParticipant]:
    head = None
    for r in range(1, 40):
        cols = _header_columns(sheet, r, range(1, 20))
        if all(h in cols for h in ANMELDUNG_HEADERS):
            head = (r, cols)
            break
    if head is None:
        _structure(issues, program, "Anmeldung: Kopfzeile der Teilnehmertabelle nicht gefunden.")
        return []
    row, cols = head
    people = []
    for r in range(row + 1, sheet.max_row + 1):
        excel_id = sheet.get(r, cols["ID"])
        if isinstance(excel_id, bool) or not isinstance(excel_id, (int, float)):
            continue  # leere oder unvollständige Anmeldezeile
        values = {h: sheet.get(r, cols[h]) for h in ANMELDUNG_HEADERS}
        people.append(RawParticipant(
            excel_id=int(excel_id), row=r, name=values["Name"], vorname=values["Vorname"], jahrgang=values["Jg."],
            verein=values["Verein"], ak=values["AK"], ak_kontr=values["akKontr"], startgeraet=values["Startgeraet"],
        ))
    return people


def _read_testprogramm(sheet: _Sheet, issues: Issues, program: str) -> tuple[list[LineDefinition], dict]:
    head = None
    for r in range(1, 30):
        cols = _header_columns(sheet, r, range(1, 20))
        if all(h in cols for h in ("Number", "Column1", "VersatzMin", "VersatzMax")):
            head = (r, cols)
            break
    lines = []
    if head is None:
        _structure(issues, program, "Testprogramm: Kopfzeile der Tabelle testprogramm nicht gefunden.")
    else:
        row, cols = head
        for r in range(row + 1, sheet.max_row + 1):
            number = sheet.raw(r, cols["Number"])
            match = LINE_NUMBER.match(number.strip()) if isinstance(number, str) else None
            if not match:
                continue
            sheet.get(r, cols["Number"])
            name = sheet.get(r, cols["Column1"])
            vmin, vmax = sheet.get(r, cols["VersatzMin"]), sheet.get(r, cols["VersatzMax"])
            lines.append(LineDefinition(
                geraet_code=int(match.group(1)), nummer=int(match.group(2)), name=clean_text(name),
                versatz_min=int(vmin) if isinstance(vmin, (int, float)) and not isinstance(vmin, bool) else None,
                versatz_max=int(vmax) if isinstance(vmax, (int, float)) and not isinstance(vmax, bool) else None,
            ))
    minmax = {}
    ak_head = None
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_col + 1):
            if sheet.text(r, c) == "AK" and sheet.text(r, c + 1) == "min" and sheet.text(r, c + 2) == "max":
                ak_head = (r, c)
                break
        if ak_head:
            break
    if ak_head is None:
        _structure(issues, program, "Testprogramm: Tabelle akMinMax (Kopf 'AK', 'min', 'max') nicht gefunden.")
    else:
        row, col = ak_head
        last = col
        while sheet.text(row, last + 1):
            last += 1
        for r in range(row + 1, sheet.max_row + 1):
            ak = sheet.raw(r, col)
            if isinstance(ak, bool) or not isinstance(ak, (int, float)):
                if ak is None:
                    break
                continue
            minmax[int(ak)] = tuple(sheet.get(r, c) for c in range(col, last + 1))
    return lines, minmax


def _read_code_legend(sheet: _Sheet) -> Optional[dict[str, str]]:
    """Legende der Einteilungscodes (`id_athlete` … `id_position`, darunter die Zeichen `@ ° ç ^`) oder None."""
    for r in range(1, min(sheet.max_row, 30) + 1):
        for c in range(1, sheet.max_col + 1):
            if sheet.text(r, c) != EINTEILUNG_FELDER[0]:
                continue
            cols = {sheet.text(r, cc): cc for cc in range(c, min(c + 10, sheet.max_col + 1))}
            if all(f in cols for f in EINTEILUNG_FELDER):
                return {f: sheet.text(r + 1, cols[f]) for f in EINTEILUNG_FELDER}
    return None


def _read_startliste(sheet: _Sheet, stations: list[Station], issues: Issues,
                     program: str) -> tuple[list[StartlistEntry], list[str], Optional[dict[str, str]]]:
    """Abschnitte je Startgerät mit Position und ID; je Person zusätzlich die Einteilungscodes der Rotationsspalten
    (`@ID°Rotation çGerät ^Startposition`). Eine Rotationsspalte hat im Abschnittskopf die Rotationsnummer und in der
    Zeile darunter den Gerätenamen."""
    names = {s.name for s in stations}
    entries, sections = [], []
    section, id_col, rot_cols = None, None, []
    for r in range(1, sheet.max_row + 1):
        first = sheet.text(r, 1)
        header_id = next((c for c in range(2, sheet.max_col + 1) if sheet.text(r, c) == "ID"), None)
        if first in names and header_id is not None:
            section, id_col = first, header_id
            sections.append(first)
            rot_cols = []
            for c in range(header_id + 1, sheet.max_col + 1):
                nummer = sheet.raw(r, c)
                if (isinstance(nummer, (int, float)) and not isinstance(nummer, bool)
                        and sheet.text(r + 1, c) in names):
                    rot_cols.append((c, sheet.get(r, c), sheet.text(r + 1, c)))
            continue
        position = sheet.raw(r, 1)
        if section is None or isinstance(position, bool) or not isinstance(position, (int, float)):
            continue
        excel_id = sheet.get(r, id_col)
        position = sheet.get(r, 1)
        if isinstance(excel_id, bool) or not isinstance(excel_id, (int, float)):
            _structure(issues, program, f"Startliste!{col_letter(id_col)}{r}: interne ID fehlt.")
            continue
        codes = tuple((nummer, geraet, sheet.get(r, c), f"Startliste!{col_letter(c)}{r}")
                      for c, nummer, geraet in rot_cols)
        entries.append(StartlistEntry(section, position, int(excel_id), r, codes))
    if not sections:
        _structure(issues, program, "Startliste: keine Abschnitte je Startgerät gefunden.")
    return entries, sections, _read_code_legend(sheet)


def display_text(value: Any) -> str:
    """Zeitplanwerte als Darstellungstext – ohne Uhrzeitberechnung."""
    if value is None:
        return ""
    if isinstance(value, dtime):
        return value.strftime("%H:%M:%S" if value.second else "%H:%M")
    if isinstance(value, datetime):
        if value.year < 1901:  # Excel-Uhrzeit ohne Datum
            return value.strftime("%H:%M:%S" if value.second else "%H:%M")
        if value.hour == value.minute == value.second == 0:
            return value.strftime("%d.%m.%Y")
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def _read_zeitplan(sheet: _Sheet, issues: Issues, program: str) -> list[dict]:
    cols = _header_columns(sheet, 1, range(1, sheet.max_col + 1))
    if "Zeit" not in cols or "Programmpunkt" not in cols:
        _structure(issues, program, "Zeitplan: Kopfzeile mit 'Zeit' und 'Programmpunkt' fehlt.")
        return []
    keys = [("zeit", cols["Zeit"]), ("programmpunkt", cols["Programmpunkt"])]
    if "Hinweis" in cols:
        keys.append(("hinweis", cols["Hinweis"]))
    rows = []
    for r in range(2, sheet.max_row + 2):
        values = {key: display_text(sheet.get(r, c)) for key, c in keys}
        if not any(values.values()):
            break  # erste vollständig leere Zeile beendet den Zeitplan
        values.setdefault("hinweis", "")
        rows.append(values)
    return rows


def _read_noteneingabe(sheet: _Sheet, stations: list[Station], issues: Issues, program: str) -> dict:
    out = {"noten_rows": [], "slots": [], "direct_fields": [], "onsite_blocks": [], "technik_blocks": []}
    head = None
    for r in range(1, 20):
        for c in range(1, 30):
            if all(sheet.text(r + i, c) == label for i, label in enumerate(NOTEN_LABELS)):
                head = (r, c)
                break
        if head:
            break
    if head is None:
        _structure(issues, program, "Noteneingabe: Kopfzeilen (Gerät, Linie, Sonderzeichen, Suchbereich, "
                                    "Ergebnissbereich) nicht gefunden.")
        return out
    r_geraet, label_col = head
    r_linie, r_sonder, r_such, r_ergebnis, r_feld = (r_geraet + i for i in range(1, 6))
    id_head = sheet.find("ID", range(r_feld + 1, r_feld + 6), range(1, label_col + 1))
    if id_head is None:
        _structure(issues, program, "Noteneingabe: Kopfzeile mit 'ID' nicht gefunden.")
        return out
    id_row, id_col = id_head
    for r in range(id_row + 1, sheet.max_row + 1):
        value = sheet.raw(r, id_col)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            sheet.get(r, id_col)
            out["noten_rows"].append((int(value), r))

    technik = {s.name for s in stations if s.name != ATHLETIK_STATION}
    stations_all = {s.name for s in stations}
    blocks = []  # (bereich, geraet, nummer, columns{feld: col})
    c = label_col + 1
    while c <= sheet.max_col:
        feld = sheet.text(r_feld, c)
        geraet = sheet.text(r_geraet, c)
        marker = f"{sheet.text(r_linie, c)} {sheet.text(r_such, c)}"
        if feld == BLOCK_FIELDS[0]:
            cols = {}
            for i, expected in enumerate(BLOCK_FIELDS):
                cc = c + i
                if (sheet.text(r_feld, cc) != expected or sheet.text(r_geraet, cc) != geraet
                        or sheet.raw(r_linie, cc) != sheet.raw(r_linie, c)):
                    _structure(issues, program, f"Noteneingabe!{col_letter(c)}: unvollständiger 8-Spalten-Block.")
                    cols = None
                    break
                sheet.get(r_geraet, cc)
                sheet.get(r_linie, cc)
                cols[expected] = cc
            if cols is None:
                c += 1
                continue
            nummer = sheet.get(r_linie, c)
            if isinstance(nummer, bool) or not isinstance(nummer, (int, float)) or int(nummer) != nummer:
                _structure(issues, program, f"Noteneingabe!{col_letter(c)}{r_linie}: Linie ist keine ganze Zahl.")
                c += len(BLOCK_FIELDS)
                continue
            nummer = int(nummer)
            kraft = KRAFT_SHEET.match(geraet)
            if geraet in technik:
                blocks.append((BEREICH_TECHNIK, geraet, nummer, cols))
                out["technik_blocks"].append((geraet, nummer, col_letter(c)))
            elif geraet == ATHLETIK_STATION and geraet in stations_all:
                blocks.append((BEREICH_BEWEGLICHKEIT, ATHLETIK_STATION, nummer, cols))
                out["onsite_blocks"].append(OnsiteBlock(BEREICH_BEWEGLICHKEIT, nummer, geraet, col_letter(c)))
            elif kraft and geraet.startswith(f"9{nummer}"):
                blocks.append((BEREICH_KRAFT, "Kraft", nummer, cols))
                out["onsite_blocks"].append(OnsiteBlock(BEREICH_KRAFT, nummer, geraet, col_letter(c)))
            else:
                _structure(issues, program, f"Noteneingabe!{col_letter(c)}: Block mit unbekanntem Gerät oder "
                                            f"unpassender Testnummer.")
            c += len(BLOCK_FIELDS)
            continue
        if "total" in marker:
            c += 1
            continue
        if geraet in ("Beweglichkeit", "Kraft") and feld == "End" and not sheet.text(r_ergebnis, c) and sheet.text(r_sonder, c):
            nummer = sheet.get(r_linie, c)
            if isinstance(nummer, bool) or not isinstance(nummer, (int, float)) or int(nummer) != nummer:
                _structure(issues, program, f"Noteneingabe!{col_letter(c)}{r_linie}: Testnummer ist keine ganze Zahl.")
            else:
                bereich = BEREICH_BEWEGLICHKEIT if geraet == "Beweglichkeit" else BEREICH_KRAFT
                sheet.get(r_sonder, c)
                out["direct_fields"].append(DirectField(bereich, int(nummer), sheet.text(r_sonder, c), col_letter(c)))
                blocks.append((bereich, geraet, int(nummer), {"End": c, "_direkt": True}))
        c += 1

    for excel_id, row in out["noten_rows"]:
        for bereich, geraet, nummer, cols in blocks:
            if cols.get("_direkt"):
                out["slots"].append(ScoreSlot(excel_id, bereich, QUELLE_DIREKT, geraet, nummer,
                                              end=sheet.get(row, cols["End"]),
                                              zelle=f"Noteneingabe!{col_letter(cols['End'])}{row}"))
                continue
            athlete = sheet.get(row, cols["id_athlete"])
            if isinstance(athlete, (int, float)) and not isinstance(athlete, bool) and int(athlete) != excel_id:
                issues.add(Code.ERGEBNIS_ID_DOPPELT, f"{program}: Noteneingabe!{col_letter(cols['id_athlete'])}{row}: "
                                                     "Ergebnis gehört nicht zur Person dieser Zeile.")
            quelle = QUELLE_STATION
            out["slots"].append(ScoreSlot(
                excel_id, bereich, quelle, geraet, nummer,
                rotation=sheet.get(row, cols["id_rotation"]), sks=sheet.get(row, cols["SKS"]),
                tal=sheet.get(row, cols["TAL"]), end=sheet.get(row, cols["End"]),
                zelle=f"Noteneingabe!{col_letter(cols['End'])}{row}",
            ))
    return out


# ── Öffentliche Funktionen ───────────────────────────────────────────────────────────────────────────

def read_program_source(
    path: Path, *, wait_seconds: float = 2.0, sleep: Callable[[float], None] = time.sleep
) -> ProgramSource:
    """Liest eine Arbeitsmappe (über eine temporäre Kopie) – die Quelle bleibt unverändert."""
    path = Path(path)
    issues = Issues()
    with excel_probe.stable_temp_copy(path, wait_seconds, sleep) as (copy_path, size, _mtime_ns):
        source = SourceInfo(filename=path.name, sha256=excel_probe.sha256_file(copy_path), size_bytes=size)
        workbook = excel_probe._load(copy_path, data_only=True)
        try:
            missing = [name for name in SHEETS if name not in workbook.sheetnames]
            if missing:
                issues.add(Code.STRUKTUR_UNGUELTIG, f"{path.name}: Pflichtblätter fehlen: {', '.join(missing)}.")
                issues.raise_if_any()
            flags = scan_cell_flags(copy_path, SHEETS)
            sheets = {}
            label = path.name
            for name in SHEETS:
                sheets[name] = _Sheet(name, workbook[name], flags[name], issues, label)
            public = _read_public(workbook["Public"], sheets["Public"], issues)
            label = public.get("program") or path.name
            for sheet in sheets.values():
                sheet.program = label
            stations = _read_stations(sheets["Anmeldung"], issues, label)
            participants = _read_participants(sheets["Anmeldung"], issues, label)
            lines, minmax = _read_testprogramm(sheets["Testprogramm"], issues, label)
            startliste, sections, zeichen = _read_startliste(sheets["Startliste"], stations, issues, label)
            zeitplan = _read_zeitplan(sheets["Zeitplan"], issues, label)
            noten = _read_noteneingabe(sheets["Noteneingabe"], stations, issues, label)
        finally:
            workbook.close()
    issues.raise_if_any()
    return ProgramSource(
        public=public, source=source, stations=stations, line_definitions=lines, ak_minmax=minmax,
        participants=participants, startliste=startliste, startliste_sections=sections, zeitplan=zeitplan,
        einteilung_zeichen=zeichen, **noten,
    )


def load_publication_state(
    paths: Iterable[Path], *, observed_at, previous: Optional[dict] = None, wait_seconds: float = 2.0,
    sleep: Callable[[float], None] = time.sleep,
) -> dict:
    """Liest alle Arbeitsmappen eines Testtags und erzeugt den kanonischen Publikationsstand."""
    sources = [read_program_source(Path(p), wait_seconds=wait_seconds, sleep=sleep) for p in paths]
    return build_state(sources, observed_at=observed_at, previous=previous)
