"""Lesende Excelprüfung für P0.

Die Quelldatei wird nie geöffnet, um sie zu verändern: Nach einer Stabilitätsprüfung wird eine temporäre
Kopie angelegt, ausschliesslich diese Kopie wird mit openpyxl gelesen und danach wieder entfernt.

openpyxl berechnet keine Formeln. Gelesen werden nur die von Excel gespeicherten Resultatwerte
(`data_only=True`) und – getrennt davon – die Formeln selbst (`data_only=False`).

Das erzeugte Diagnoseobjekt enthält nur Metadaten, Zählwerte und Prüfsummen, keine Zellinhalte aus
`Noteneingabe`, `Startliste` oder `Zeitplan`.
"""
from __future__ import annotations

import contextlib
import hashlib
import json
import os
import re
import shutil
import stat as stat_module
import tempfile
import time
import warnings
import zipfile
from datetime import date, datetime, time as dtime, timedelta, timezone
from pathlib import Path
from typing import Callable, Iterator

SCHEMA_VERSION = "probe-1"
REQUIRED_SHEETS = ("Public", "Noteneingabe", "Startliste", "Zeitplan")
# (Zelle Wert, Zelle Beschriftung, Feldname)
PUBLIC_FIELDS = (
    ("B2", "A2", "event_id"),
    ("B3", "A3", "event_title"),
    ("B4", "A4", "program"),
    ("B5", "A5", "published_rotation_max"),
    ("B6", "A6", "public_active"),
)
EMPTY_SAMPLE_LIMIT = 20
EVENT_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,99}$")

# Windows-Dateiattribute für nicht vollständig lokal vorhandene OneDrive-Dateien (Platzhalter).
_FILE_ATTRIBUTE_OFFLINE = 0x1000
_FILE_ATTRIBUTE_RECALL_ON_OPEN = 0x40000
_FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS = 0x400000
_PLACEHOLDER_ATTRIBUTES = (
    _FILE_ATTRIBUTE_OFFLINE | _FILE_ATTRIBUTE_RECALL_ON_OPEN | _FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS
)


class ProbeError(Exception):
    """Basisfehler mit Exitcode für die Kommandozeile."""

    exit_code = 1


class SourceError(ProbeError):
    """Exceldatei fehlt, ist gesperrt, instabil, nur teilweise synchronisiert oder ungültig."""

    exit_code = 3


class ValidationError(ProbeError):
    """Arbeitsmappe ist lesbar, entspricht aber nicht der erwarteten Struktur."""

    exit_code = 4


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _utc_iso(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _snapshot(path: Path) -> tuple[int, int]:
    try:
        st = path.stat()
    except FileNotFoundError:
        raise SourceError(f"Exceldatei nicht gefunden: {path.name}") from None
    except OSError as exc:
        raise SourceError(f"Exceldatei nicht lesbar ({type(exc).__name__}): {path.name}") from None
    if not stat_module.S_ISREG(st.st_mode):
        raise SourceError(f"Keine reguläre Datei: {path.name}")
    attributes = getattr(st, "st_file_attributes", 0)
    if attributes & _PLACEHOLDER_ATTRIBUTES:
        raise SourceError(
            f"Datei ist nicht vollständig lokal synchronisiert (OneDrive-Platzhalter): {path.name}. "
            "Im Explorer 'Immer auf diesem Gerät behalten' wählen und Synchronisation abwarten."
        )
    return st.st_size, st.st_mtime_ns


def check_source_path(path: Path) -> None:
    name = path.name
    if name.startswith("~$"):
        raise SourceError(f"Excel-Sperrdatei statt Arbeitsmappe angegeben: {name}")
    if path.suffix.lower() != ".xlsx":
        raise SourceError(f"Nur .xlsx-Arbeitsmappen werden unterstützt: {name}")


def lock_file_present(path: Path) -> bool:
    """Excel legt beim Öffnen `~$<name>` an. Hinweis, aber kein Abbruchgrund."""
    return (path.parent / ("~$" + path.name[2:])).exists() or (path.parent / ("~$" + path.name)).exists()


def wait_until_stable(
    path: Path, wait_seconds: float, sleep: Callable[[float], None] = time.sleep
) -> tuple[int, int]:
    """Grösse und Änderungszeit zweimal erfassen; nur bei Gleichheit weiter."""
    first = _snapshot(path)
    if first[0] == 0:
        raise SourceError(f"Exceldatei ist leer (Synchronisation unvollständig?): {path.name}")
    sleep(wait_seconds)
    second = _snapshot(path)
    if first != second:
        raise SourceError(
            f"Exceldatei verändert sich noch (Speichern/Synchronisation läuft): {path.name}. Später erneut versuchen."
        )
    return second


def _open_shared(path: Path):
    """Quelle nur lesend öffnen, ohne andere Programme zu behindern.

    Excel (AutoSpeichern, OneDrive/SharePoint) hält die geöffnete Arbeitsmappe mit Lösch-/Umbenennungszugriff
    offen. Windows erlaubt dann nur Leser, die FILE_SHARE_DELETE zulassen – `open()` tut das nicht und scheitert
    mit einer Freigabeverletzung. Gelesen wird der zuletzt gespeicherte Stand; die Datei wird nie beschrieben.
    """
    if os.name != "nt":
        return open(path, "rb")
    import ctypes
    import msvcrt
    from ctypes import wintypes

    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CreateFileW.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                   wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE)
    # GENERIC_READ; FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE; OPEN_EXISTING; FILE_ATTRIBUTE_NORMAL
    handle = kernel.CreateFileW(str(path), 0x80000000, 0x7, None, 3, 0x80, None)
    if handle in (None, wintypes.HANDLE(-1).value):
        raise ctypes.WinError(ctypes.get_last_error())  # 32/33 → PermissionError, 2/3 → FileNotFoundError
    try:
        return os.fdopen(msvcrt.open_osfhandle(handle, os.O_RDONLY | os.O_BINARY), "rb")
    except BaseException:
        kernel.CloseHandle(handle)
        raise


@contextlib.contextmanager
def stable_temp_copy(
    path: Path, wait_seconds: float, sleep: Callable[[float], None] = time.sleep
) -> Iterator[tuple[Path, int, int]]:
    """Liefert (Kopie, Grösse, mtime_ns) einer stabilen Quelldatei; die Kopie wird danach entfernt."""
    check_source_path(path)
    size, mtime_ns = wait_until_stable(path, wait_seconds, sleep)
    with tempfile.TemporaryDirectory(prefix="testtag-probe-") as tmp:
        copy_path = Path(tmp) / "source.xlsx"
        try:
            with _open_shared(path) as src, open(copy_path, "wb") as dst:
                shutil.copyfileobj(src, dst, 1024 * 1024)
        except PermissionError:
            raise SourceError(f"Exceldatei ist gesperrt und kann nicht gelesen werden: {path.name}") from None
        except OSError as exc:
            raise SourceError(f"Exceldatei konnte nicht kopiert werden ({type(exc).__name__}): {path.name}") from None
        if _snapshot(path) != (size, mtime_ns) or copy_path.stat().st_size != size:
            raise SourceError(f"Exceldatei hat sich während des Kopierens verändert: {path.name}")
        if not zipfile.is_zipfile(copy_path):
            raise SourceError(f"Keine gültige .xlsx-Datei (kein ZIP-Container): {path.name}")
        yield copy_path, size, mtime_ns


def _load(path: Path, data_only: bool):
    import openpyxl

    with warnings.catch_warnings():
        # openpyxl warnt bei nicht unterstützten Erweiterungen (z. B. Datenüberprüfung); für P0 irrelevant.
        warnings.simplefilter("ignore", UserWarning)
        try:
            return openpyxl.load_workbook(
                path, read_only=True, data_only=data_only, keep_vba=False, keep_links=False
            )
        except Exception as exc:  # beschädigte oder fremde Datei
            raise SourceError(f"Arbeitsmappe konnte nicht gelesen werden ({type(exc).__name__}).") from None


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and value == "")


def _normalize_cached(value) -> str | float | bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, (datetime, date, dtime)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if value is None:
        return None
    return str(value)


def _iter_content_cells(ws) -> Iterator:
    """Alle gespeicherten Zellen mit Inhalt – unabhängig von der deklarierten Dimension."""
    for row in ws.iter_rows():
        for cell in row:
            if getattr(cell, "value", None) is not None:
                yield cell


def _range_string(min_row, min_col, max_row, max_col) -> str:
    from openpyxl.utils import get_column_letter

    if min_row is None:
        return ""
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def analyse_noteneingabe(ws_formulas, ws_cached) -> dict:
    from openpyxl.utils import get_column_letter

    try:
        declared = ws_formulas.calculate_dimension()
    except ValueError:
        declared = ""
    ws_formulas.reset_dimensions()
    ws_cached.reset_dimensions()

    min_row = min_col = max_row = max_col = None
    formula_coords: list[tuple[int, int]] = []
    for cell in _iter_content_cells(ws_formulas):
        r, c = cell.row, cell.column
        min_row = r if min_row is None else min(min_row, r)
        max_row = r if max_row is None else max(max_row, r)
        min_col = c if min_col is None else min(min_col, c)
        max_col = c if max_col is None else max(max_col, c)
        if cell.data_type == "f":
            formula_coords.append((r, c))

    cached: dict[tuple[int, int], tuple[str, object]] = {}
    for cell in _iter_content_cells(ws_cached):
        cached[(cell.row, cell.column)] = (cell.data_type, cell.value)

    digest = hashlib.sha256()
    nonempty = empty = 0
    empty_samples: list[str] = []
    for r, c in sorted(formula_coords):
        coord = f"{get_column_letter(c)}{r}"
        data_type, value = cached.get((r, c), ("n", None))
        if _is_empty(value):
            empty += 1
            if len(empty_samples) < EMPTY_SAMPLE_LIMIT:
                empty_samples.append(coord)
            record = [coord, "empty", None]
        else:
            nonempty += 1
            record = [coord, data_type, _normalize_cached(value)]
        digest.update(json.dumps(record, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        digest.update(b"\n")

    return {
        "dimension": _range_string(min_row, min_col, max_row, max_col),
        "dimension_declared": declared,
        "formula_cells": len(formula_coords),
        "cached_nonempty_cells": nonempty,
        "cached_empty_cells": empty,
        "cached_values_sha256": digest.hexdigest(),
        "cached_empty_sample_addresses": empty_samples,
    }


def _cell_value(ws, coordinate: str):
    return ws[coordinate].value


def read_public(ws) -> dict:
    raw = {}
    for value_cell, label_cell, field in PUBLIC_FIELDS:
        label = _cell_value(ws, label_cell)
        if not isinstance(label, str) or label.strip() != field:
            raise ValidationError(f"Public!{label_cell} muss die Beschriftung '{field}' enthalten.")
        raw[field] = _cell_value(ws, value_cell)

    event_id = raw["event_id"]
    if not isinstance(event_id, str) or not EVENT_ID_PATTERN.match(event_id.strip()):
        raise ValidationError("Public!B2 (event_id) fehlt oder enthält unzulässige Zeichen (erlaubt: A-Z a-z 0-9 . _ -).")
    for coordinate, field in (("B3", "event_title"), ("B4", "program")):
        value = raw[field]
        if not isinstance(value, str) or not value.strip():
            raise ValidationError(f"Public!{coordinate} ({field}) fehlt oder ist kein Text.")

    rotation = raw["published_rotation_max"]
    if isinstance(rotation, float) and rotation.is_integer():
        rotation = int(rotation)
    if isinstance(rotation, bool) or not isinstance(rotation, int) or not 0 <= rotation <= 99:
        raise ValidationError("Public!B5 (published_rotation_max) muss eine ganze Zahl zwischen 0 und 99 sein.")

    active = raw["public_active"]
    if not isinstance(active, bool):
        raise ValidationError("Public!B6 (public_active) muss WAHR oder FALSCH sein.")

    return {
        "event_id": event_id.strip(),
        "event_title": raw["event_title"].strip(),
        "program": raw["program"].strip(),
        "published_rotation_max": rotation,
        "public_active": active,
    }


def count_zeitplan_rows(ws) -> int:
    """Anzahl nicht leerer Zeilen unterhalb der Kopfzeile (Zeile 1)."""
    ws.reset_dimensions()
    rows = set()
    for cell in _iter_content_cells(ws):
        if cell.row > 1 and not (isinstance(cell.value, str) and not cell.value.strip()):
            rows.add(cell.row)
    return len(rows)


def probe_workbook(
    path: Path,
    *,
    wait_seconds: float = 2.0,
    now: datetime | None = None,
    sleep: Callable[[float], None] = time.sleep,
) -> dict:
    """Liest die Arbeitsmappe (über eine temporäre Kopie) und liefert das Diagnoseobjekt."""
    path = Path(path)
    generated_at = now or datetime.now(timezone.utc)
    with stable_temp_copy(path, wait_seconds, sleep) as (copy_path, size, mtime_ns):
        source_sha = sha256_file(copy_path)
        wb_f = _load(copy_path, data_only=False)
        try:
            wb_v = _load(copy_path, data_only=True)
            try:
                sheets = {name: name in wb_f.sheetnames for name in REQUIRED_SHEETS}
                missing = [name for name, ok in sheets.items() if not ok]
                if missing:
                    raise ValidationError("Pflichtblätter fehlen: " + ", ".join(missing))
                event = read_public(wb_v["Public"])
                noten = analyse_noteneingabe(wb_f["Noteneingabe"], wb_v["Noteneingabe"])
                zeitplan_rows = count_zeitplan_rows(wb_v["Zeitplan"])
            finally:
                wb_v.close()
        finally:
            wb_f.close()

    modified = datetime.fromtimestamp(mtime_ns / 1e9, tz=timezone.utc)
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": _utc_iso(generated_at),
        "event": event,
        "source": {
            "filename": path.name,
            "modified_at": _utc_iso(modified),
            "size_bytes": size,
            "sha256": source_sha,
        },
        "sheets": sheets,
        "noteneingabe": noten,
        "zeitplan": {"data_row_count": zeitplan_rows},
    }


def serialize_report(report: dict) -> bytes:
    """Stabile UTF-8-Serialisierung (Feldreihenfolge wie aufgebaut, LF, abschliessender Zeilenumbruch)."""
    return (json.dumps(report, ensure_ascii=False, indent=2) + "\n").encode("utf-8")

