"""Gemeinsame Hilfen für die Tests: Importpfad und synthetische Arbeitsmappen (nur temporäre Verzeichnisse)."""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True
_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import socket  # noqa: E402

import openpyxl  # noqa: E402


def _no_network(*_args, **_kwargs):
    raise AssertionError("Unerwarteter echter Netzwerkzugriff im Unit-Test")


# Globale Sperre: Kein Test darf eine echte Verbindung aufbauen.
socket.socket.connect = _no_network
socket.socket.connect_ex = _no_network
socket.create_connection = _no_network
socket.getaddrinfo = _no_network

# Synthetische, aber realistisch aussehende Teilnehmerdaten. Keiner dieser Werte darf im JSON auftauchen.
PARTICIPANTS = [
    ("Muster", "Anna", 2014, "TV Synthetikon", "P-4711", "R-07", 9.35),
    ("Beispiel", "Ben", 2015, "STV Fiktiv", "P-4712", "R-08", 8.125),
]
SECRET_MARKERS = ["Muster", "Anna", "Beispiel", "Ben", "2014", "2015", "Synthetikon", "Fiktiv",
                  "P-4711", "P-4712", "R-07", "R-08", "9.35", "8.125", "18.7", "16.25", "Turnhalle Nord"]

DEFAULT_PUBLIC = {
    "event_id": "synth-2026-ak5-ak9plus",
    "event_title": "Synthetischer Testtag AK5–AK9+",
    "program": "AK5-AK9+",
    "published_rotation_max": 7,
    "public_active": True,
}


def make_workbook(
    path: Path,
    *,
    sheets=("Public", "Noteneingabe", "Startliste", "Zeitplan"),
    public: dict | None = None,
    cached: dict | None = None,
    extra_noten_cell: str | None = None,
) -> Path:
    """Erzeugt eine synthetische Arbeitsmappe. `cached` setzt gespeicherte Formelresultate (Adresse -> Wert)."""
    public = dict(DEFAULT_PUBLIC, **(public or {}))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        if name == "Public":
            ws["A1"] = "Public-Konfiguration"
            for row, key in enumerate(DEFAULT_PUBLIC, start=2):
                ws.cell(row, 1, key)
                ws.cell(row, 2, public[key])
        elif name == "Noteneingabe":
            ws.append(["Name", "Vorname", "Note", "Doppelt", "Bonus", "ID"])
            for i, (name_, first, _y, _v, pid, _r, score) in enumerate(PARTICIPANTS, start=2):
                ws.append([name_, first, score, f"=C{i}*2", f'=IF(C{i}>9,1,"")', pid])
            if extra_noten_cell:
                ws[extra_noten_cell] = "=C2+1"
        elif name == "Startliste":
            ws.append(["Name", "Vorname", "Jahrgang", "Verein", "Riege"])
            for name_, first, year, club, _p, riege, _s in PARTICIPANTS:
                ws.append([name_, first, year, club, riege])
        elif name == "Zeitplan":
            ws.append(["Zeit", "Programmpunkt", "Hinweis"])
            ws.append(["07:30", "Hallenöffnung", "Turnhalle Nord"])
            ws.append(["08:00", "Rotation 1", None])
            ws.append([None, None, None])
            ws.append(["09:00", "Rotation 2", None])
    wb.save(path)
    if cached is None:
        cached = {"D2": 18.7, "D3": 16.25, "E2": 1, "E3": ""}
    inject_cached(path, sheets.index("Noteneingabe") + 1 if "Noteneingabe" in sheets else None, cached)
    return path


def inject_cached(path: Path, sheet_number: int | None, cached: dict) -> None:
    """openpyxl schreibt Formeln ohne Resultat; hier werden gespeicherte Resultate wie von Excel ergänzt."""
    if sheet_number is None or not cached:
        return
    member = f"xl/worksheets/sheet{sheet_number}.xml"
    with zipfile.ZipFile(path) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    out = []
    for info, data in items:
        if info.filename == member:
            xml = data.decode("utf-8")
            for coord, value in cached.items():
                pattern = re.compile(rf'<c r="{coord}"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>')
                if isinstance(value, str):
                    repl = rf'<c r="{coord}" t="str"><f>\2</f><v>{value}</v></c>'
                elif isinstance(value, bool):
                    repl = rf'<c r="{coord}" t="b"><f>\2</f><v>{int(value)}</v></c>'
                else:
                    repl = rf'<c r="{coord}"><f>\2</f><v>{value}</v></c>'
                xml, n = pattern.subn(repl, xml)
                assert n == 1, f"Formelzelle {coord} nicht gefunden"
            data = xml.encode("utf-8")
        out.append((info, data))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in out:
            zout.writestr(info, data)


# ── P1B: synthetische Testtag-Arbeitsmappen (Struktur wie die produktiven Mappen, erfundene Personen) ──

import copy as _copy  # noqa: E402
from datetime import time as _time  # noqa: E402

TT_SLUG = "synth-testtag-2026-10-09"
TT_TITLE = "Synthetischer Testtag 2026 – Musterhausen"
AK5 = "AK5-AK9+"
AK9 = "AK9-AK12"

AK5_BEW = ["Bücke", "Arm-Rumpf-Winkel rücklings", "Ein- und Ausschultern", "Brücke", "Briefmarke",
           "Querspagat", "Seitspagat"]
AK5_KRAFT = ["Standweitsprung", "Klettern / Hangeln am Tau", "Handstand", "Kraftkomplex am Barren / Winkelstütz",
             "Liegestütz"]
AK9_BEW = ["Rumpfbeugen vorwärts", "Arm-Rumpf-Winkel rücklings", "Ein- und Ausschultern", "Brücke", "Querspagat",
           "Seitspagat", "Arm-Rumpf-Winkel vorlings", "Beinspreizen vorwärts", "Beinspreizen seitwärts",
           "Fussgelenkwinkel"]
AK9_KRAFT = ["Sprint", "Dreierhupf", "Hangeln", "Klappmesser", "Kreisflanken", "Handstand", "Schweizer",
             "Wiener/Pless", "Handstandschieber", "Winkel – Spitzwinkel", "Hangwaage rücklings", "Hangwaage vorlings",
             "Stützwaage", "Klimmzug-Beugestütz-Zugstemme"]

# Erfundene Personen; keiner dieser Werte darf in Fehlermeldungen oder Prüflauf-Ausgaben erscheinen.
TT_SECRET_MARKERS = ["Synthaler", "Fiktivo", "Probst", "Imaginaer", "Kunstmann", "Erdacht", "Modell", "Vorlage",
                     "Lea", "Nina", "Noah", "Elias", "Testdorf", "Beispielstadt"]


def _ak5_spec() -> dict:
    return {
        "public": {"event_id": TT_SLUG, "event_title": TT_TITLE, "program": AK5, "published_rotation_max": 8,
                   "public_active": True, "Kraftmodus": "vor_ort", "Beweglichkeitsmodus": "vor_ort"},
        "stations": [("Boden", 1), ("Pferd", 2), ("Ringe", 3), ("Beweglichkeit", 4), ("Sprung", 5), ("Barren", 6),
                     ("Trampolin", 7), ("Reck", 8)],
        "aks": [(5, "AK 5"), (6, "AK 6"), (7, "AK 7"), (8, "AK 8"), (9, "AK9+")],
        # (Excel-ID, Name, Vorname, Jahrgang, Verein, AK, akKontr, Startgerät)
        "participants": [
            (1, "Synthaler", "Lea", 2018, "TV Testdorf", "AK 7", 7, "Boden"),
            (2, "Fiktivo", "Max", 2017, "TV Testdorf", "AK 7", 7, "Pferd"),
            (3, "Probst", "Nina", 2019, "STV Beispielstadt", "AK 6", 6, "Pferd"),
            (4, "Imaginaer", "Tim", 2018, "STV Beispielstadt", "AK 7", 7, "Beweglichkeit"),
        ],
        # (Gerät, Linie, Name, {akKontr: (min, max)} – fehlende AK: (1, 10))
        "technik_lines": [
            ("Boden", 1, "Synth Boden vorwärts", {}), ("Boden", 2, "Synth Boden rückwärts", {6: (0, 0)}),
            ("Pferd", 1, "Synth Scheren", {}), ("Ringe", 1, "Synth Schwünge", {}), ("Sprung", 1, "Synth Überschlag", {}),
            ("Barren", 1, "Synth Stütz", {}), ("Trampolin", 1, "Synth Salto", {}), ("Reck", 1, "Synth Felgen", {}),
        ],
        "extra_blocks": [("Boden", 3)],  # Block in Noteneingabe ohne Testprogramm-Definition
        "bew_names": AK5_BEW, "kraft_names": AK5_KRAFT, "onsite": True,
        "kraft_station": {1: "Ringe", 2: "Sprung", 3: "Reck", 4: "Barren", 5: "Pferd"},
        "kraft_sheet": {1: "91Standweitsprung", 2: "92Klettern", 3: "93Handstand", 4: "94Winkel", 5: "95Liegestütz"},
        "zeitplan": [(_time(7, 30), "Hallenöffnung", "AK5–AK9+"), ("08:05–12:30", "Wettkampf", None),
                     (None, None, None), ("nach Leerzeile", "nicht übernehmen", None)],
        "scores": {}, "direct": {}, "rotation": {}, "cells": {},
    }


def _ak9_spec() -> dict:
    return {
        "public": {"event_id": TT_SLUG, "event_title": TT_TITLE, "program": AK9, "published_rotation_max": 7,
                   "public_active": True, "Kraftmodus": "extern", "Beweglichkeitsmodus": "extern"},
        "stations": [("Boden", 1), ("Pferd", 2), ("Ringe", 3), ("Sprung", 5), ("Barren", 6), ("Trampolin", 7),
                     ("Reck", 8)],
        "aks": [(9, "AK 9"), (10, "AK 10"), (11, "AK 11"), (12, "AK 12")],
        "participants": [
            (1, "Kunstmann", "Jan", 2016, "TV Testdorf", "AK 9", 9, "Boden"),
            (2, "Erdacht", "Luca", 2016, "STV Beispielstadt", "AK 9", 9, "Sprung"),
            (3, "Modell", "Noah", 2015, "TV Testdorf", "AK 10", 10, "Sprung"),
            (4, "Vorlage", "Elias", 2014, "STV Beispielstadt", "AK 11", 11, "Reck"),
        ],
        "technik_lines": [
            ("Boden", 1, "Synth Akrobatik", {}), ("Pferd", 1, "Synth Scheren", {}), ("Ringe", 1, "Synth Schwünge", {}),
            ("Sprung", 1, "Überschlag vw", {}), ("Sprung", 2, "Tsukahara", {9: (0, 0)}),
            ("Sprung", 3, "Yurchenko", {9: (0, 0)}), ("Barren", 1, "Synth Stütz", {}),
            ("Trampolin", 1, "Synth Salto", {}), ("Reck", 1, "Synth Felgen", {}),
        ],
        "extra_blocks": [],
        "bew_names": AK9_BEW, "kraft_names": AK9_KRAFT, "onsite": False, "kraft_station": {}, "kraft_sheet": {},
        "zeitplan": [(_time(13, 55), "Besammlung und Einmarsch", "AK9–AK12"), ("14:00 – 16:00", "Wettkampf", None)],
        "scores": {}, "direct": {}, "rotation": {}, "cells": {},
    }


def testtag_spec(program: str = AK5, **changes) -> dict:
    """Standardspezifikation je Programm; `public`, `scores`, `direct`, `rotation`, `cells` werden ergänzt."""
    spec = _ak5_spec() if program == AK5 else _ak9_spec()
    for key, value in changes.items():
        if key in ("public", "scores", "direct", "rotation", "cells"):
            spec[key] = dict(spec[key], **value) if key == "public" else {**spec[key], **value}
        else:
            spec[key] = _copy.deepcopy(value)
    return spec


def expected_rotation(spec: dict, start: str, station: str) -> int:
    order = [name for name, _code in spec["stations"]]
    return (order.index(station) - order.index(start)) % len(order) + 1


def make_testtag_workbook(path: Path, program: str = AK5, spec: dict | None = None, **changes) -> Path:
    """Synthetische Arbeitsmappe mit der Struktur der produktiven Testtag-Mappen.

    `scores[(id, gerät, nummer)] = (SKS, TAL, End)` – Gerät ist der Stationsname, 'Beweglichkeit' oder 'Kraft'
    (Vor-Ort-Blöcke). `direct[(id, 'Beweglichkeit'|'Kraft', nummer)] = End`. Eine Zeichenkette '=…' ergibt eine
    Formel ohne gespeichertes Resultat, '#N/A' einen Excel-Fehlerwert.
    """
    from openpyxl.utils import get_column_letter

    spec = spec or testtag_spec(program, **changes)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    station_names = [n for n, _c in spec["stations"]]
    codes = dict(spec["stations"])

    ws = wb.create_sheet("Public")
    ws["A1"] = "Public-Konfiguration"
    for row, key in enumerate(("event_id", "event_title", "program", "published_rotation_max", "public_active",
                               "Kraftmodus", "Beweglichkeitsmodus"), start=2):
        ws.cell(row, 1, key)
        ws.cell(row, 2, spec["public"].get(key))
    if "subdivision" in spec["public"]:  # optional: Subdivision eines Programms (Public!A9/B9)
        ws.cell(9, 1, "subdivision")
        ws.cell(9, 2, spec["public"]["subdivision"])

    ws = wb.create_sheet("Anmeldung")
    for col, head in enumerate(("ID", "Name", "Vorname", "Jg.", "Verein", "AK", "Startgeraet", "GeraetID", "akKontr"), 1):
        ws.cell(9, col, head)
    for row, (pid, name, first, year, club, ak, kontr, start) in enumerate(spec["participants"], start=10):
        for col, value in enumerate((pid, name, first, year, club, ak, start, codes.get(start), kontr), 1):
            ws.cell(row, col, value)
    for col, head in enumerate(("startgeraet", "Column1", "GerID"), 11):
        ws.cell(9, col, head)
    ws.cell(9, 15, "AK")  # Tab_AK: zweite Kopfzelle 'AK' in derselben Zeile wie die Teilnehmertabelle
    for row, (_kontr, label) in enumerate(spec["aks"], start=10):
        ws.cell(row, 15, label)
    for row, (name, code) in enumerate(spec["stations"], start=10):
        ws.cell(row, 11, name)
        ws.cell(row, 12, row - 9)
        ws.cell(row, 13, code)

    ws = wb.create_sheet("Testprogramm")
    for col, head in enumerate(("Gerät", "Linie", "Stufe", "SKS", "Number", "Column1", "min", "max", "VersatzMin",
                                "VersatzMax"), 1):
        ws.cell(2, col, head)
    row, versatz = 3, 2
    minmax_cols = {}
    for name, code in spec["stations"]:
        if name == "Beweglichkeit":
            continue
        ws.cell(row, 5, f"ç{code}§0|0"); ws.cell(row, 6, name); row += 1
        for g, linie, line_name, limits in spec["technik_lines"]:
            if g != name:
                continue
            ws.cell(row, 5, f"ç{code}§{linie}|0"); ws.cell(row, 6, line_name)
            ws.cell(row, 9, versatz); ws.cell(row, 10, versatz + 1)
            minmax_cols[(g, linie)] = (versatz, limits)
            row += 1
            ws.cell(row, 5, f"ç{code}§{linie}|1"); ws.cell(row, 6, "Stufe 1"); row += 1
            versatz += 2
    for n, test in enumerate(spec["bew_names"], 1):
        ws.cell(row, 5, f"ç4§{n}|0"); ws.cell(row, 6, test); row += 1
    ws.cell(row, 5, "ç9§0|0"); ws.cell(row, 6, "Kraft"); row += 1
    for n, test in enumerate(spec["kraft_names"], 1):
        ws.cell(row, 5, f"ç9§{n}|0"); ws.cell(row, 6, test); row += 1
    ws.cell(15, 12, "AK"); ws.cell(15, 13, "min"); ws.cell(15, 14, "max")
    for v in range(4, versatz + 1):
        ws.cell(15, 11 + v, f"{'min' if v % 2 == 0 else 'max'}{v}")
    for r, (kontr, _label) in enumerate(spec["aks"], start=16):
        ws.cell(r, 12, kontr)
        for (_g, _l), (v, limits) in minmax_cols.items():
            lo, hi = limits.get(kontr, (1, 10))
            ws.cell(r, 11 + v, lo)
            ws.cell(r, 12 + v, hi)

    # Startliste wie produktiv: Legende der Einteilungscodes, je Abschnitt Rotationskopf (L…) mit Gerätezeile darunter,
    # je Person der Code @ID°Rotation çGerät ^Startposition (ç = Position in tab_startgeraet), dazu eine weitere
    # Zahlenspalte ohne Gerät. `startliste_codes[(id, rotation)]` ersetzt einen Code, `ohne_legende` lässt sie weg.
    ws = wb.create_sheet("Startliste")
    ws["A1"] = "Synthetischer Testtag"
    if not spec.get("ohne_legende"):
        for col, (feld, zeichen) in enumerate(zip(("id_athlete", "id_rotation", "id_apparatus", "id_position"),
                                                  ("@", "°", "ç", "^")), 13):
            ws.cell(3, col, feld); ws.cell(4, col, zeichen)
    ersatz = spec.get("startliste_codes", {})
    n = len(station_names)
    row = 6
    for name in station_names:
        ws.cell(row, 1, name); ws.cell(row, 8, "ID")
        for col, head in enumerate(("Pos", "Name", "Vorname", "Jahrgang", "Kategorie", "Verein"), 1):
            ws.cell(row + 1, col, head)
        folge = [station_names[(station_names.index(name) + r) % n] for r in range(n)]
        for r, geraet in enumerate(folge, 1):
            ws.cell(row, 11 + r, r); ws.cell(row + 1, 11 + r, geraet)
        ws.cell(row, 12 + n, n + 1); ws.cell(row + 1, 12 + n, "Total")
        row += 2
        group = [p for p in spec["participants"] if p[7] == name]
        for pos, (pid, pname, first, year, club, ak, _k, _s) in enumerate(group, 1):
            for col, value in enumerate((pos, pname, first, year, ak, club), 1):
                ws.cell(row, col, value)
            ws.cell(row, 8, pid)
            for r, geraet in enumerate(folge, 1):
                code = f"@{pid}°{r}ç{station_names.index(geraet) + 1}^{pos}"
                ws.cell(row, 11 + r, ersatz.get((pid, r), code))
            row += 1
        row += 2

    ws = wb.create_sheet("Zeitplan")
    for col, head in enumerate(("Zeit", "Programmpunkt", "Hinweis"), 1):
        ws.cell(1, col, head)
    for r, values in enumerate(spec["zeitplan"], start=2):
        for col, value in enumerate(values, 1):
            ws.cell(r, col, value)

    ws = wb.create_sheet("Noteneingabe")
    ws["A1"] = "Noteneingabe"
    for r, label in enumerate(("Gerät", "Linie", "Sonderzeichen", "Suchbereich", "Ergebnissbereich"), start=4):
        ws.cell(r, 6, label)
    for col, head in enumerate(("ID", "Name", "Vorname", "Jg.", "Verein", "AK"), 1):
        ws.cell(10, col, head)
    people = sorted(spec["participants"], key=lambda p: (station_names.index(p[7]), p[0]))
    rows, r = {}, 11
    for name in station_names:
        group = [p for p in people if p[7] == name]
        if not group:
            continue
        ws.cell(r, 1, name); r += 1
        for p in group:
            for col, value in enumerate((p[0], p[1], p[2], p[3], p[4], p[5]), 1):
                ws.cell(r, col, value)
            rows[p[0]] = (r, p)
            r += 1

    col = 7
    marks = ("@", "°", "ç", "^", "§", None, None, None)
    fields = ("id_athlete", "id_rotation", "id_apparatus", "id_position", "id_linie", "SKS", "TAL", "End")

    def block(label, nummer, kind, station, score_key):
        nonlocal col
        for i, field in enumerate(fields):
            c = col + i
            ws.cell(4, c, label); ws.cell(5, c, nummer); ws.cell(6, c, marks[i])
            ws.cell(7, c, "AF14:AF2000"); ws.cell(8, c, f"{get_column_letter(27 + i)}14:{get_column_letter(27 + i)}2000")
            ws.cell(9, c, field)
        for pid, (rr, p) in rows.items():
            rot = spec["rotation"].get((pid, score_key, nummer), expected_rotation(spec, p[7], station))
            sks, tal, end = spec["scores"].get((pid, score_key, nummer), (None, 1 if kind == "athletik" else None, None))
            values = (pid, rot, station_names.index(station) + 1, 1, nummer, sks, tal, end)
            for i, value in enumerate(values):
                ws.cell(rr, col + i, value)
        col += len(fields)

    def direct(bereich, names):
        nonlocal col
        for n, test in enumerate(names, 1):
            ws.cell(4, col, bereich); ws.cell(5, col, n); ws.cell(6, col, test)
            ws.cell(7, col, "keine Übernahmeformeln"); ws.cell(9, col, "End")
            for pid, (rr, _p) in rows.items():
                ws.cell(rr, col, spec["direct"].get((pid, bereich, n)))
            col += 1

    for g, linie, _n, _l in spec["technik_lines"]:
        block(g, linie, "technik", g, g)
    for g, linie in spec["extra_blocks"]:
        block(g, linie, "technik", g, g)
    if spec["onsite"]:
        for n in range(1, len(spec["bew_names"]) + 1):
            block("Beweglichkeit", n, "athletik", "Beweglichkeit", "Beweglichkeit")
    direct("Beweglichkeit", spec["bew_names"])
    if spec["onsite"]:
        for n in range(1, len(spec["kraft_names"]) + 1):
            block(spec["kraft_sheet"][n], n, "athletik", spec["kraft_station"][n], "Kraft")
    direct("Kraft", spec["kraft_names"])
    # Excel-Totale mit absichtlich falschen Werten: dürfen nie übernommen werden.
    for label, total in (("Boden", "total_boden_end"), ("Kraft", "total_kraft_end")):
        ws.cell(4, col, label); ws.cell(5, col, total); ws.cell(9, col, "End")
        for _pid, (rr, _p) in rows.items():
            ws.cell(rr, col, 999)
        col += 1
    ws.cell(4, col, "TOTAL"); ws.cell(7, col, "total_athletik_technik_end"); ws.cell(9, col, "End")

    for (sheet, coord), value in spec["cells"].items():
        wb[sheet][coord] = value
    wb.save(path)
    return path


# ── Simulierte Supabase Data API (PostgREST) – vollständig im Speicher ─────────────────────────────────

import json as _json  # noqa: E402
import urllib.parse as _parse  # noqa: E402
import uuid as _uuid  # noqa: E402

TEST_SECRET_KEY = "sb_secret_TESTONLY_q1w2e3r4t5y6u7i8o9p0"
TEST_PUBLISHABLE_KEY = "sb_publishable_TESTONLY_a1s2d3f4g5h6j7k8"
TEST_ENV = {
    "SUPABASE_URL": "https://example-ref.supabase.co",
    "SUPABASE_SECRET_KEY": TEST_SECRET_KEY,
    "SUPABASE_PUBLISHABLE_KEY": TEST_PUBLISHABLE_KEY,
}
VEROEFFENTLICHUNG_PK = ("oeffentliche_testtag_id", "bereich", "programm", "ak")


def jsonb_reorder(value):
    """Ahmt JSONB nach: Objektfelder kommen in anderer Reihenfolge zurück (Länge, dann Bytes)."""
    if isinstance(value, dict):
        return {k: jsonb_reorder(value[k]) for k in sorted(value, key=lambda k: (len(k.encode()), k))}
    if isinstance(value, list):
        return [jsonb_reorder(v) for v in value]
    return value


class FakeSupabase:
    def __init__(self, *, secure=True, tamper_read=False, read_rows=None, write_status=None):
        self.secure = secure
        self.tamper_read = tamper_read
        self.read_rows = read_rows
        self.write_status = write_status
        self.testtage = {}
        self.veroeffentlichung = {}
        self.calls = []

    @staticmethod
    def _json(status, value):
        return status, _json.dumps(value).encode("utf-8")

    def __call__(self, method, url, headers, body, timeout):
        self.calls.append({"method": method, "url": url, "headers": dict(headers),
                           "body": None if body is None else _json.loads(body)})
        parts = _parse.urlsplit(url)
        table = parts.path.rsplit("/", 1)[-1]
        params = _parse.parse_qsl(parts.query, keep_blank_values=True)
        key = headers.get("apikey")
        if key == TEST_SECRET_KEY:
            role = "service_role"
        elif key == TEST_PUBLISHABLE_KEY:
            role = "anon"
        else:
            return self._json(401, {"message": "Invalid API key"})
        payload = None if body is None else _json.loads(body)

        if role == "anon" and method in ("POST", "PATCH", "DELETE"):
            if self.secure:
                return self._json(401, {"code": "42501", "message": f"permission denied for table {table}"})
            return 201, b""
        if role == "service_role" and method == "POST":
            if self.write_status:
                return self._json(self.write_status, {"code": "XX000", "message": f"boom {key}"})
            if table == "testtage":
                existing = self.testtage.get(payload["slug"], {"id": str(_uuid.uuid4()), "status": "vorbereitet"})
                existing.update(payload)
                self.testtage[payload["slug"]] = existing
                return self._json(201, [{"id": existing["id"]}])
            if table == "veroeffentlichung":
                pk = tuple(payload[k] for k in VEROEFFENTLICHUNG_PK)
                self.veroeffentlichung[pk] = payload
                return 201, b""
        if method == "GET" and table == "veroeffentlichung":
            if self.read_rows is not None:
                return self._json(200, self.read_rows)
            filters = {k: v[3:] for k, v in params if k != "select" and v.startswith("eq.")}
            public_ids = {t["id"] for t in self.testtage.values() if t.get("ist_oeffentlich")}
            rows = []
            for row in self.veroeffentlichung.values():
                if role == "anon" and row["oeffentliche_testtag_id"] not in public_ids:
                    continue
                if all(str(row[k]) == v for k, v in filters.items()):
                    out = dict(row)
                    out["inhalt"] = jsonb_reorder(row["inhalt"])
                    out["stand_am"] = row["stand_am"].replace("Z", "+00:00")
                    if self.tamper_read:
                        out["inhalt"] = dict(out["inhalt"], schema_version="manipuliert")
                    rows.append(out)
            return self._json(200, rows)
        return self._json(400, {"message": "nicht simuliert"})


# ── P1C: simulierte Data API mit RLS, Rechten, Fremdschlüssel und atomarem Batch-Upsert ─────────────────

_MANAGED = ("zeitplan", "startliste", "resultate_live", "resultate_rangliste")
_PK = ("oeffentliche_testtag_id", "bereich", "programm", "ak")
_HEX64 = re.compile(r"^[0-9a-f]{64}$")


def _db_time(value: str) -> str:
    return value.replace("Z", "+00:00")


class FakeDataApi:
    """PostgREST-Teilmenge für `testtage` und `veroeffentlichung`, vollständig im Speicher.

    Fehlerinjektion: `fail_batch` (HTTP-Status), `batch_drop` (Antwort ohne eine Zeile),
    `tamper_public` (Funktion, die öffentlich gelesene Zeilen verändert), `fail_patch` (HTTP-Status).
    """

    def __init__(self, *, fail_batch=None, batch_drop=False, tamper_public=None, fail_patch=None):
        self.testtage: dict[str, dict] = {}
        self.rows: dict[tuple, dict] = {}
        self.snapshots: dict[str, dict] = {}
        self.calls: list[dict] = []
        self.fail_batch, self.batch_drop = fail_batch, batch_drop
        self.tamper_public, self.fail_patch = tamper_public, fail_patch
        self.fail_archive_status = None  # HTTP-Status für PATCH status=archiviert (Abbruch nach dem Snapshot)
        self.offline = False  # True: Netzwerkunterbruch (TransportError)
        self.clock = 0

    # Hilfen für Tests
    def add_testtag(self, slug, *, titel=TT_TITLE, status="live", ist_oeffentlich=True, stand_am="2026-10-01T00:00:00Z"):
        tid = str(_uuid.uuid4())
        self.testtage[tid] = {"id": tid, "slug": slug, "titel": titel, "datum": None, "ort": None, "status": status,
                              "ist_oeffentlich": ist_oeffentlich, "stand_am": _db_time(stand_am)}
        return tid

    def add_row(self, tid, bereich, programm, ak, inhalt, *, stand_am="2026-10-01T00:00:00Z", pruefsumme=None):
        from testtag_public.publication_model import content_checksum
        self.rows[(tid, bereich, programm, ak)] = {
            "oeffentliche_testtag_id": tid, "bereich": bereich, "programm": programm, "ak": ak,
            "pruefsumme": pruefsumme or content_checksum(inhalt), "stand_am": _db_time(stand_am), "inhalt": inhalt}

    def writes(self):
        return [c for c in self.calls if c["method"] in ("POST", "PATCH", "DELETE")]

    @staticmethod
    def _json(status, value):
        return status, _json.dumps(value).encode("utf-8")

    @staticmethod
    def _match(row, filters):
        for key, op in filters:
            if op.startswith("eq."):
                if str(row.get(key)).lower() != op[3:].lower() if key in ("id", "oeffentliche_testtag_id") \
                        else str(row.get(key)) != op[3:]:
                    return False
            elif op.startswith("in.("):
                if str(row.get(key)) not in op[4:-1].split(","):
                    return False
        return True

    @staticmethod
    def _select(row, params):
        fields = dict(params).get("select")
        return {k: row[k] for k in fields.split(",")} if fields else dict(row)

    def __call__(self, method, url, headers, body, timeout):
        if self.offline:
            from testtag_public.supabase_database import TransportError
            raise TransportError("URLError: getaddrinfo failed")
        parts = _parse.urlsplit(url)
        table = parts.path.rsplit("/", 1)[-1]
        params = _parse.parse_qsl(parts.query, keep_blank_values=True)
        payload = None if body is None else _json.loads(body)
        self.calls.append({"method": method, "table": table, "params": params, "headers": dict(headers),
                           "body": payload})
        key = headers.get("apikey")
        role = {TEST_SECRET_KEY: "service_role", TEST_PUBLISHABLE_KEY: "anon"}.get(key)
        if role is None:
            return self._json(401, {"message": "Invalid API key"})
        filters = [(k, v) for k, v in params if k not in ("select", "on_conflict")]
        public_ids = {t["id"] for t in self.testtage.values() if t["ist_oeffentlich"]}
        if method != "GET" and role == "anon":
            return self._json(401, {"code": "42501", "message": f"permission denied for table {table}"})
        if method == "DELETE":
            return self._json(403, {"code": "42501", "message": f"permission denied for table {table}"})

        if table == "testtage":
            if method == "GET":
                rows = [t for t in self.testtage.values() if self._match(t, filters)
                        and (role == "service_role" or t["ist_oeffentlich"])]
                return self._json(200, [self._select(t, params) for t in rows])
            if method == "POST":
                if any(t["slug"] == payload["slug"] for t in self.testtage.values()):
                    return self._json(409, {"code": "23505", "message": "duplicate key value violates unique constraint"})
                tid = self.add_testtag(payload["slug"], titel=payload["titel"], status=payload.get("status", "vorbereitet"),
                                       ist_oeffentlich=payload.get("ist_oeffentlich", False),
                                       stand_am=f"2026-01-01T00:00:{self.clock:02d}Z")
                return self._json(201, [self._select(self.testtage[tid], params)])
            if method == "PATCH":
                if self.fail_patch:
                    return self._json(self.fail_patch, {"code": "XX000", "message": "boom"})
                if self.fail_archive_status and payload.get("status") == "archiviert":
                    return self._json(self.fail_archive_status, {"code": "XX000", "message": "boom"})
                rows = [t for t in self.testtage.values() if self._match(t, filters)]
                for t in rows:
                    t.update({k: (_db_time(v) if k == "stand_am" else v) for k, v in payload.items()})
                return self._json(200, [self._select(t, params) for t in rows])

        if table == "veroeffentlichung":
            if method == "GET":
                out = []
                for row in self.rows.values():
                    if not self._match(row, filters):
                        continue
                    if role == "anon" and row["oeffentliche_testtag_id"] not in public_ids:
                        continue
                    out.append(self._select(dict(row, inhalt=jsonb_reorder(row["inhalt"])), params))
                if role == "anon" and self.tamper_public:
                    out = self.tamper_public(out)
                return self._json(200, out)
            if method == "POST":
                if self.fail_batch:
                    # Serverantwort mit Zeileninhalt: darf nie in einer Meldung erscheinen.
                    return self._json(self.fail_batch, {"code": "23514", "message":
                                      f"Failing row contains ({_json.dumps(payload, ensure_ascii=False)[:300]})"})
                items = payload if isinstance(payload, list) else [payload]
                for item in items:  # atomar: erst alles prüfen
                    if item["oeffentliche_testtag_id"] not in self.testtage:
                        return self._json(409, {"code": "23503", "message": "violates foreign key constraint"})
                    if item["bereich"] not in (*_MANAGED, "probe") or not _HEX64.match(item["pruefsumme"]):
                        return self._json(400, {"code": "23514", "message": "violates check constraint"})
                for item in items:
                    self.rows[tuple(item[k] for k in _PK)] = {**item, "stand_am": _db_time(item["stand_am"])}
                returned = [self._select(self.rows[tuple(i[k] for k in _PK)], params) for i in items]
                if self.batch_drop:
                    returned = returned[:-1]
                return self._json(201, returned)

        if table == "archiv_snapshots":
            # Rechte wie schema.sql: anon nur SELECT (RLS: öffentlich und archiviert), service_role SELECT/INSERT.
            if method == "GET":
                archived_public = {t["id"] for t in self.testtage.values()
                                   if t["ist_oeffentlich"] and t["status"] == "archiviert"}
                out = [self._select(dict(s, inhalt=jsonb_reorder(s["inhalt"])), params)
                       for s in self.snapshots.values() if self._match(s, filters)
                       and (role == "service_role" or s["oeffentliche_testtag_id"] in archived_public)]
                return self._json(200, out)
            if method == "POST":
                if payload["oeffentliche_testtag_id"] not in self.testtage:
                    return self._json(409, {"code": "23503", "message": "violates foreign key constraint"})
                if not _HEX64.match(payload["pruefsumme"]):
                    return self._json(400, {"code": "23514", "message": "violates check constraint"})
                if payload["oeffentliche_testtag_id"] in self.snapshots:
                    return self._json(409, {"code": "23505", "message": "duplicate key value"})
                row = {**payload, "eingefroren_am": f"2026-01-01T00:00:{self.clock:02d}+00:00"}
                self.snapshots[payload["oeffentliche_testtag_id"]] = row
                return self._json(201, [self._select(row, params)])
            return self._json(403, {"code": "42501", "message": f"permission denied for table {table}"})
        return self._json(400, {"message": "nicht simuliert"})


# ── P1D: synthetische öffentliche Publikation (echtes P1B/P1C-Format) und öffentliche Lese-Schnittstelle ──

PUB_T1 = "2026-10-09T08:00:00Z"
PUB_T2 = "2026-10-09T09:30:00Z"
LIVE_ID = "11111111-1111-4111-8111-111111111111"
ARCHIV_ID = "22222222-2222-4222-8222-222222222222"
VORBEREITET_ID = "33333333-3333-4333-8333-333333333333"
PRIVAT_ID = "44444444-4444-4444-8444-444444444444"
ARCHIV_SLUG = "synth-testtag-2025-10-10"
PUB_SECRET_MARKERS = [*TT_SECRET_MARKERS, TEST_SECRET_KEY, TEST_PUBLISHABLE_KEY, "Failing row", "Traceback",
                      "TAL", "excel_id", "ergebnis_id", "sb_secret_"]


def synthetic_publication(directory: Path) -> dict:
    """Erzeugt Testtage und `veroeffentlichung`-Zeilen mit dem echten P1B/P1C-Code (nur erfundene Personen)."""
    from testtag_public import publication_model as pm
    from testtag_public.database_publisher import tombstone_inhalt
    from testtag_public.excel_publication import read_program_source

    ak5_people = testtag_spec(AK5)["participants"] + [(5, "Ersonnen", "Mia", 2018, "TV Testdorf", "AK 7", 7, "Reck")]
    first = {(1, "Boden", 1): (5, 1.5, 7.5), (2, "Boden", 1): (6, 1, 6), (1, "Kraft", 1): (None, 1, 8)}
    later = {**first, (4, "Pferd", 1): (6, 1, 6)}  # neu in der zweiten Beobachtung
    ak9_scores = {(1, "Sprung", 1): (3, 1, 3), (1, "Sprung", 2): (4, 1, 4), (3, "Sprung", 1): (5, 1, 5),
                  (3, "Sprung", 3): (6, 1, 6)}

    def sources(ak5_scores):
        paths = [make_testtag_workbook(directory / f"pub-{_uuid.uuid4().hex}.xlsx", AK5,
                                       participants=ak5_people, scores=ak5_scores),
                 make_testtag_workbook(directory / f"pub-{_uuid.uuid4().hex}.xlsx", AK9,
                                       scores=ak9_scores, direct={(3, "Kraft", 2): 4})]
        return [read_program_source(p, wait_seconds=0, sleep=lambda _s: None) for p in paths]

    state1 = pm.build_state(sources(first), observed_at=PUB_T1)
    state2 = pm.build_state(sources(later), observed_at=PUB_T2, previous=state1)

    def rows_for(testtag_id, state, stand_am):
        rows = [{"oeffentliche_testtag_id": testtag_id, "bereich": r["bereich"], "programm": r["programm"],
                 "ak": r["ak"], "pruefsumme": r["pruefsumme"], "stand_am": _db_time(stand_am), "inhalt": r["inhalt"]}
                for r in pm.publication_rows(state)]
        tomb = tombstone_inhalt("resultate_live")
        rows.append({"oeffentliche_testtag_id": testtag_id, "bereich": "resultate_live", "programm": AK5, "ak": "AK 5",
                     "pruefsumme": pm.content_checksum(tomb), "stand_am": _db_time(stand_am), "inhalt": tomb})
        rows.append({"oeffentliche_testtag_id": testtag_id, "bereich": "probe", "programm": AK5, "ak": "",
                     "pruefsumme": "0" * 64, "stand_am": _db_time(stand_am),
                     "inhalt": {"schema_version": "probe-1", "event": {"program": "PROBEINHALT"}}})
        return rows

    testtage = [
        {"id": LIVE_ID, "slug": TT_SLUG, "titel": TT_TITLE, "status": "live", "stand_am": _db_time(PUB_T2),
         "ist_oeffentlich": True, "datum": None, "ort": None},
        {"id": ARCHIV_ID, "slug": ARCHIV_SLUG, "titel": "Synthetischer Testtag 2025 – Archivhausen",
         "status": "archiviert", "stand_am": _db_time("2025-10-10T16:00:00Z"), "ist_oeffentlich": True,
         "datum": None, "ort": None},
        {"id": VORBEREITET_ID, "slug": "synth-testtag-2027", "titel": "Synthetischer Testtag 2027",
         "status": "vorbereitet", "stand_am": _db_time("2026-09-01T10:00:00Z"), "ist_oeffentlich": True,
         "datum": None, "ort": None},
        {"id": PRIVAT_ID, "slug": "synth-geheim", "titel": "Nicht öffentlich", "status": "live",
         "stand_am": _db_time(PUB_T2), "ist_oeffentlich": False, "datum": None, "ort": None},
    ]
    archiv_rows = rows_for(ARCHIV_ID, state1, "2025-10-10T16:00:00Z")
    rows = rows_for(LIVE_ID, state2, PUB_T2) + archiv_rows
    rows += rows_for(PRIVAT_ID, state1, PUB_T1)
    from testtag_public.archive import ARCHIV_SCHEMA_VERSION, archive_checksum, build_archive
    archiv = build_archive(testtage[1], archiv_rows)
    snapshots = [{"oeffentliche_testtag_id": ARCHIV_ID, "schema_version": ARCHIV_SCHEMA_VERSION,
                  "pruefsumme": archive_checksum(archiv), "inhalt": archiv}]
    return {"testtage": testtage, "rows": rows, "snapshots": snapshots, "state": state2}


class FakeOeffentlich:
    """Öffentliche PostgREST-Leseschnittstelle (Rolle anon, RLS) für `quelle.py`: transport(url, kopf, zeit)."""

    def __init__(self, testtage, rows, snapshots=()):
        self.testtage = _copy.deepcopy(testtage)
        self.rows = _copy.deepcopy(rows)
        self.snapshots = _copy.deepcopy(list(snapshots))
        self.calls: list[dict] = []
        self.status = None  # z. B. 503: HTTP-Fehler mit Inhalt in der Antwort
        self.ausnahme = None  # Exception, die der Transport wirft
        self.manipulieren = None  # Funktion(rows) -> rows für veroeffentlichung
        self.manipulieren_archiv = None  # Funktion(snapshots) -> snapshots für archiv_snapshots

    def __call__(self, url, kopf, zeitueberschreitung):
        self.calls.append({"url": url, "kopf": dict(kopf), "zeit": zeitueberschreitung})
        if self.ausnahme is not None:
            raise self.ausnahme
        if self.status:
            return self.status, _json.dumps({"message": "Failing row contains (Synthaler, TV Testdorf)",
                                              "hint": TEST_PUBLISHABLE_KEY})
        if kopf.get("apikey") != TEST_PUBLISHABLE_KEY:
            return 401, '{"message":"Invalid API key"}'
        parts = _parse.urlsplit(url)
        table = parts.path.rsplit("/", 1)[-1]
        params = _parse.parse_qsl(parts.query, keep_blank_values=True)
        select = dict(params).get("select", "")
        filters = [(k, v[3:]) for k, v in params if k != "select" and v.startswith("eq.")]
        public = {t["id"] for t in self.testtage if t["ist_oeffentlich"]}
        if table == "testtage":
            data = [t for t in self.testtage if t["ist_oeffentlich"] and all(str(t.get(k)) == v for k, v in filters)]
        elif table == "veroeffentlichung":
            data = [r for r in self.rows if r["oeffentliche_testtag_id"] in public
                    and all(str(r.get(k)) == v for k, v in filters)]
            data = [dict(r, inhalt=jsonb_reorder(r["inhalt"])) for r in data]
            if self.manipulieren:
                data = self.manipulieren(_copy.deepcopy(data))
        elif table == "archiv_snapshots":
            archiviert = {t["id"] for t in self.testtage if t["ist_oeffentlich"] and t["status"] == "archiviert"}
            data = [dict(s, inhalt=jsonb_reorder(s["inhalt"])) for s in self.snapshots
                    if s["oeffentliche_testtag_id"] in archiviert and all(str(s.get(k)) == v for k, v in filters)]
            if self.manipulieren_archiv:
                data = self.manipulieren_archiv(_copy.deepcopy(data))
        else:
            return 404, ""
        fields = select.split(",") if select else None
        return 200, _json.dumps([{k: d.get(k) for k in fields} if fields else d for d in data], ensure_ascii=False)
