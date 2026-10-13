"""Synthetische Publikation im 30-s-Takt für den Lasttest (aus dem Streamlit-Lasttest übernommen).

    python -B publikationstakt.py <projekt> <datenordner> <sekunden> <publog>
<datenordner> enthält voll_slug.txt (systemtest-last-…) und voll/ak5.xlsx, voll/ak9.xlsx (synthetisch).
Schreibschutz: nur der systemtest-last-Slug, nur verwaltete Bereiche. Protokolliert je Publikation den Zeitpunkt
der Supabase-Bestätigung (Antwort auf das PATCH von testtage.stand_am).
"""
import io
import json
import os
import random
import sys
import time
import urllib.parse
import uuid
from pathlib import Path

PROJEKT, SCRATCH, ARGS = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve(), sys.argv[3:]
sys.path.insert(0, str(PROJEKT / "src"))
sys.dont_write_bytecode = True


def takt():
    import openpyxl
    from testtag_public import local_config as lc
    from testtag_public import publisher as cli
    from testtag_public import supabase_database as db
    from testtag_public.database_publisher import MANAGED_BEREICHE
    config = db.load_config(lc.load_env(os.environ))
    slug = (SCRATCH / "voll_slug.txt").read_text().strip()
    assert slug.startswith("systemtest-last-")
    dauer, log = float(ARGS[0]), Path(ARGS[1])
    ak5, ak9 = SCRATCH / "voll" / "ak5.xlsx", SCRATCH / "voll" / "ak9.xlsx"
    ids, marke = set(), {}

    def guard(method, url, headers, body, timeout):
        parts = urllib.parse.urlsplit(url)
        table = parts.path.rsplit("/", 1)[-1]
        assert headers.get("apikey") in (config.secret_key, config.publishable_key)
        assert not any(k.lower() == "authorization" for k in headers)
        if method != "GET":
            data = json.loads(body)
            rows = data if isinstance(data, list) else [data]
            filters = [(k, v) for k, v in urllib.parse.parse_qsl(parts.query) if k not in ("select", "on_conflict")]
            ok = (table == "testtage" and method == "PATCH" and len(filters) == 1 and filters[0][1][3:] in ids) or (
                table == "veroeffentlichung" and method == "POST" and all(
                    r.get("oeffentliche_testtag_id") in ids and r.get("bereich") in MANAGED_BEREICHE for r in rows))
            if not ok:
                raise db.TransportError("SCHUTZ")
        status, payload = db.urllib_transport(method, url, headers, body, timeout)
        if table == "testtage" and status in (200, 201):
            for r in json.loads(payload):
                if r.get("slug") == slug and r.get("id"):
                    ids.add(str(uuid.UUID(r["id"])))
        if method == "PATCH" and table == "testtage" and status == 200 and sorted(json.loads(body)) == ["stand_am"]:
            marke["supabase"] = time.time()
        return status, payload

    s, p = guard("GET", db.table_url(config, "testtage", [("select", "id,slug"), ("slug", f"eq.{slug}")]),
                 db._headers(config.secret_key), None, 20)
    rnd = random.Random(time.time_ns())
    ende, naechste = time.time() + dauer, time.time()
    while time.time() < ende:
        time.sleep(max(0.0, naechste - time.time()))
        naechste += 30
        t0 = time.time()
        wb = openpyxl.load_workbook(ak5)
        ws = wb["Noteneingabe"]
        (r,) = [r for r in range(11, ws.max_row + 1) if ws.cell(r, 1).value == 1 and ws.cell(r, 2).value]
        sp = {ws.cell(9, c).value: c for c in range(7, ws.max_column + 1)
              if ws.cell(4, c).value == "Boden" and ws.cell(5, c).value == 1}
        neu = rnd.choice([v for v in range(1, 11) if v != ws.cell(r, sp["SKS"]).value])
        ws.cell(r, sp["SKS"]).value = neu
        ws.cell(r, sp["End"]).value = neu
        wb.save(ak5)
        wb.close()
        marke.clear()
        out, err = io.StringIO(), io.StringIO()
        code = cli.run(["--workbook-ak5", str(ak5), "--workbook-ak9", str(ak9), "--stability-seconds", "0",
                        "--publish", "--confirm-slug", slug], transport=guard, stdout=out, stderr=err)
        t1 = time.time()
        with log.open("a", encoding="utf-8") as f:
            f.write(json.dumps({"t0": t0, "supabase": marke.get("supabase"), "t1": t1, "code": code,
                                "wert": neu}) + "\n")


takt()
